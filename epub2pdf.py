import os
import io
import argparse
import zipfile
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

def get_epub_info(zip_ref):
    """
    Extrai metadados do EPUB e a lista de arquivos XHTML na ordem correta do spine.
    """
    # 1. Localizar o content.opf
    container_data = zip_ref.read('META-INF/container.xml')
    soup_container = BeautifulSoup(container_data, 'xml')
    rootfile = soup_container.find('rootfile')
    if not rootfile:
        raise ValueError("container.xml inválido no EPUB.")
    opf_path = rootfile.get('full-path')
    opf_dir = os.path.dirname(opf_path)
    
    # 2. Ler o OPF
    opf_data = zip_ref.read(opf_path)
    soup_opf = BeautifulSoup(opf_data, 'xml')
    
    # Metadados
    title = soup_opf.find('dc:title') or soup_opf.find('title')
    author = soup_opf.find('dc:creator') or soup_opf.find('creator')
    
    title_text = title.text if title else "Título do Livro"
    author_text = author.text if author else "Autor do Livro"
    
    # Manifest
    manifest = soup_opf.find('manifest')
    items = {}
    if manifest:
        for item in manifest.find_all('item'):
            items[item.get('id')] = {
                'href': item.get('href'),
                'media-type': item.get('media-type')
            }
            
    # Spine
    spine = soup_opf.find('spine')
    spine_files = []
    if spine:
        for itemref in spine.find_all('itemref'):
            idref = itemref.get('idref')
            if idref in items:
                href = items[idref]['href']
                # Caminho normalizado do arquivo no zip
                full_path = os.path.normpath(os.path.join(opf_dir, href)).replace('\\', '/')
                spine_files.append((full_path, items[idref]['media-type']))
                
    return title_text, author_text, spine_files, opf_dir

def setup_section_layout(section, margin_inner_cm=2.2, margin_outer_cm=1.8, margin_top_cm=1.8, margin_bottom_cm=1.8):
    """
    Configura o tamanho A5 e margens espelhadas para a seção.
    """
    section.page_width = Cm(14.8)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(margin_top_cm)
    section.bottom_margin = Cm(margin_bottom_cm)
    section.left_margin = Cm(margin_inner_cm) # Margem interna
    section.right_margin = Cm(margin_outer_cm) # Margem externa
    
    # Habilitar margens espelhadas no XML do Word
    sectPr = section._sectPr
    mm = sectPr.find(qn('w:mirrorMargins'))
    if mm is None:
        mm = OxmlElement('w:mirrorMargins')
        sectPr.append(mm)

def add_page_number(run):
    """
    Insere o campo dinâmico PAGE no run do Word.
    """
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_table_of_contents(document):
    """
    Insere a estrutura de campo TOC do Word para o sumário automático.
    """
    p_title = document.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(18)
    p_title.paragraph_format.keep_with_next = True
    run_title = p_title.add_run("Sumário")
    run_title.font.name = 'Georgia'
    run_title.font.size = Pt(18)
    run_title.bold = True
    
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(12)
    
    run_begin = p.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    run_begin._r.append(fldChar1)
    
    run_instr = p.add_run()
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    run_instr._r.append(instrText)
    
    run_sep = p.add_run()
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    run_sep._r.append(fldChar2)
    
    run_placeholder = p.add_run("Sumário dinâmico. Pressione F9 no Word para atualizar as páginas.")
    run_placeholder.font.name = 'Georgia'
    run_placeholder.font.size = Pt(10)
    run_placeholder.font.italic = True
    
    run_end = p.add_run()
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run_end._r.append(fldChar3)

def add_title_page(document, title, author):
    """
    Cria a página de rosto do livro.
    """
    p_space = document.add_paragraph()
    p_space.paragraph_format.space_before = Pt(100)
    
    p_title = document.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.add_run(title)
    run_title.font.name = 'Georgia'
    run_title.font.size = Pt(26)
    run_title.bold = True
    
    p_sub = document.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(40)
    
    p_author = document.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_before = Pt(80)
    p_author.paragraph_format.space_after = Pt(100)
    run_author = p_author.add_run(author)
    run_author.font.name = 'Georgia'
    run_author.font.size = Pt(14)
    run_author.font.italic = True
    
    p_brand = document.add_paragraph()
    p_brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_brand.paragraph_format.space_before = Pt(50)
    run_brand = p_brand.add_run("Projeto by Jair Lima")
    run_brand.font.name = 'Georgia'
    run_brand.font.size = Pt(9)
    run_brand.font.color.rgb = RGBColor(128, 128, 128)
    
    document.add_page_break()

def add_image_to_run(img_node, paragraph, zip_ref, html_dir, max_width_cm):
    """
    Carrega a imagem do zip e a insere dimensionada na folha A5.
    """
    src = img_node.get('src')
    if not src:
        return
    
    # Resolver caminho
    img_zip_path = os.path.normpath(os.path.join(html_dir, src)).replace('\\', '/')
    
    # Tenta correspondência direta ou por nome base
    if img_zip_path not in zip_ref.namelist():
        base_name = os.path.basename(img_zip_path).lower()
        found = False
        for name in zip_ref.namelist():
            if os.path.basename(name).lower() == base_name:
                img_zip_path = name
                found = True
                break
        if not found:
            print(f"Aviso: Imagem não localizada no arquivo zip: {src}")
            return
            
    try:
        img_data = zip_ref.read(img_zip_path)
        img_stream = io.BytesIO(img_data)
        
        # Obter tamanho
        with Image.open(img_stream) as pil_img:
            width_px, height_px = pil_img.size
            dpi = pil_img.info.get('dpi', (96, 96))
            # Garantir dpi válido
            if not isinstance(dpi, tuple) or dpi[0] <= 0:
                dpi = (96, 96)
            width_cm = width_px * 2.54 / dpi[0]
            height_cm = height_px * 2.54 / dpi[1]
            
        if width_cm > max_width_cm:
            ratio = max_width_cm / width_cm
            width_cm = max_width_cm
            height_cm = height_cm * ratio
            
        img_stream.seek(0)
        
        run = paragraph.add_run()
        run.add_picture(img_stream, width=Cm(width_cm), height=Cm(height_cm))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    except Exception as e:
        print(f"Erro ao processar imagem {img_zip_path}: {e}")

def process_inline_element(node, paragraph, styles):
    """
    Processa elementos HTML em linha recursivamente para manter a formatação combinada.
    """
    current_styles = styles.copy()
    
    if node.name in ['b', 'strong']:
        current_styles['bold'] = True
    elif node.name in ['i', 'em']:
        current_styles['italic'] = True
    elif node.name == 'u':
        current_styles['underline'] = True
    elif node.name == 'sup':
        current_styles['superscript'] = True
    elif node.name == 'sub':
        current_styles['subscript'] = True
    elif node.name == 'span':
        style_attr = node.get('style', '')
        class_attr = node.get('class', [])
        if isinstance(class_attr, str):
            class_attr = [class_attr]
            
        is_red = False
        if 'red' in class_attr or any('red' in c for c in class_attr):
            is_red = True
        elif 'color' in style_attr:
            style_lower = style_attr.lower()
            if 'red' in style_lower or '#ff0000' in style_lower or '#b30000' in style_lower or 'rgb(255' in style_lower:
                is_red = True
                
        if is_red:
            current_styles['color'] = RGBColor(179, 0, 0) # Vermelho escuro premium
            
    for child in node.children:
        if isinstance(child, str):
            if child.strip() == '' and not child == ' ':
                continue
            run = paragraph.add_run(child)
            if current_styles.get('bold'):
                run.bold = True
            if current_styles.get('italic'):
                run.italic = True
            if current_styles.get('underline'):
                run.underline = True
            if current_styles.get('superscript'):
                run.font.superscript = True
            if current_styles.get('subscript'):
                run.font.subscript = True
            if current_styles.get('color'):
                run.font.color.rgb = current_styles['color']
        elif child.name == 'br':
            run = paragraph.add_run()
            run.add_break()
        elif child.name in ['b', 'strong', 'i', 'em', 'u', 'span', 'sup', 'sub', 'a']:
            process_inline_element(child, paragraph, current_styles)

def process_block(element, document, zip_ref, html_dir, max_width_cm, just_after_heading_ref):
    """
    Analisa elementos HTML de nível de bloco e insere no Word correspondente.
    """
    if element.name in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
        level = int(element.name[1])
        p = document.add_paragraph(style=f'Heading {level}')
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        
        process_inline_element(element, p, {'bold': False, 'italic': False, 'color': None, 'superscript': False, 'subscript': False, 'underline': False})
        just_after_heading_ref[0] = True
        return p
        
    elif element.name == 'p':
        # Ignorar parágrafo vazio
        if not element.text.strip() and not element.find('img'):
            return None
            
        p = document.add_paragraph()
        
        # Determinar alinhamento
        style_attr = element.get('style', '')
        class_attr = element.get('class', [])
        if isinstance(class_attr, str):
            class_attr = [class_attr]
            
        align = WD_ALIGN_PARAGRAPH.JUSTIFY
        if 'center' in style_attr or any('center' in c for c in class_attr):
            align = WD_ALIGN_PARAGRAPH.CENTER
        elif 'right' in style_attr or any('right' in c for c in class_attr):
            align = WD_ALIGN_PARAGRAPH.RIGHT
        elif 'left' in style_attr or any('left' in c for c in class_attr):
            align = WD_ALIGN_PARAGRAPH.LEFT
            
        p.alignment = align
        
        # Aplicar recuo de parágrafo (livro profissional)
        if align == WD_ALIGN_PARAGRAPH.JUSTIFY and not just_after_heading_ref[0]:
            p.paragraph_format.first_line_indent = Cm(0.5)
        else:
            p.paragraph_format.first_line_indent = Cm(0)
            
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        process_inline_element(element, p, {'bold': False, 'italic': False, 'color': None, 'superscript': False, 'subscript': False, 'underline': False})
        just_after_heading_ref[0] = False
        return p
        
    elif element.name in ['ul', 'ol']:
        for li in element.find_all('li', recursive=False):
            style_name = 'List Bullet' if element.name == 'ul' else 'List Number'
            p = document.add_paragraph(style=style_name)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            
            process_inline_element(li, p, {'bold': False, 'italic': False, 'color': None, 'superscript': False, 'subscript': False, 'underline': False})
        just_after_heading_ref[0] = True
        return None
        
    elif element.name == 'table':
        rows = element.find_all('tr')
        if not rows:
            return None
            
        num_cols = 0
        for r in rows:
            cols = r.find_all(['td', 'th'])
            num_cols = max(num_cols, len(cols))
            
        if num_cols == 0:
            return None
            
        table = document.add_table(rows=len(rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'
        
        for r_idx, r in enumerate(rows):
            cols = r.find_all(['td', 'th'])
            for c_idx, cell_node in enumerate(cols):
                if c_idx < num_cols:
                    cell = table.cell(r_idx, c_idx)
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(2)
                    
                    styles = {'bold': cell_node.name == 'th', 'italic': False, 'color': None, 'superscript': False, 'subscript': False, 'underline': False}
                    process_inline_element(cell_node, p, styles)
                    
        # Linha em branco pós tabela
        p_space = document.add_paragraph()
        p_space.paragraph_format.space_before = Pt(6)
        just_after_heading_ref[0] = True
        return table
        
    elif element.name == 'blockquote':
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.right_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        process_inline_element(element, p, {'bold': False, 'italic': True, 'color': None, 'superscript': False, 'subscript': False, 'underline': False})
        just_after_heading_ref[0] = True
        return p
        
    elif element.name == 'hr':
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run("* * *")
        run.bold = True
        just_after_heading_ref[0] = True
        return p
        
    elif element.name == 'img':
        p = document.add_paragraph()
        add_image_to_run(element, p, zip_ref, html_dir, max_width_cm)
        just_after_heading_ref[0] = True
        return p
        
    elif element.name in ['div', 'section']:
        last_p = None
        for child in element.children:
            if child.name in ['p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'ul', 'ol', 'table', 'blockquote', 'hr', 'img', 'div', 'section']:
                p_res = process_block(child, document, zip_ref, html_dir, max_width_cm, just_after_heading_ref)
                if p_res:
                    last_p = p_res
        return last_p
        
    return None

def convert_docx_to_pdf(docx_path, pdf_path):
    """
    Usa automação COM no Windows para atualizar a paginação do sumário e exportar o documento para PDF.
    """
    print("Iniciando automação do Microsoft Word via COM para geração do PDF...")
    try:
        import win32com.client
        import pythoncom
        
        pythoncom.CoInitialize()
        
        # Iniciar aplicação Word
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        
        abs_docx = os.path.abspath(docx_path)
        abs_pdf = os.path.abspath(pdf_path)
        
        print(f"Abrindo arquivo DOCX temporário: {abs_docx}")
        doc = word.Documents.Open(abs_docx)
        
        # Atualizar todos os campos
        print("Atualizando números de páginas e referências de campos...")
        doc.Fields.Update()
        
        # Garantir atualização de sumários
        print("Atualizando tabelas de conteúdo (Sumário)...")
        for toc in doc.TablesOfContents:
            toc.Update()
            
        # Salvar documento Word com o sumário paginado antes de exportar
        doc.Save()
        
        # Exportar para PDF (wdExportFormatPDF = 17)
        print(f"Exportando arquivo PDF final: {abs_pdf}")
        doc.ExportAsFixedFormat(
            OutputFileName=abs_pdf,
            ExportFormat=17,  # wdExportFormatPDF
            OpenAfterExport=False,
            OptimizeFor=0,  # wdExportOptimizeForPrint
            Range=0,  # wdExportAllDocument
            Item=0,  # wdExportDocumentContent
            IncludeDocProps=True,
            KeepIRM=True,
            CreateBookmarks=1,  # wdExportCreateHeadingBookmarks
            DocStructureTags=True,
            BitmapMissingFonts=True,
            UseISO19005_1=False
        )
        
        doc.Close()
        word.Quit()
        print("Geração do PDF concluída com sucesso pelo Microsoft Word.")
        return True
    except Exception as e:
        print(f"Erro crítico ao gerar o PDF via MS Word COM: {e}")
        print("Certifique-se de que o Microsoft Word está instalado no Windows e fechado antes de rodar.")
        return False

def main():
    parser = argparse.ArgumentParser(description="Conversor de EPUB para PDF A5 Premium via Word COM (Projeto by Jair Lima)")
    parser.add_argument("epub_path", help="Caminho do arquivo EPUB de entrada")
    parser.add_argument("pdf_path", help="Caminho do arquivo PDF de saída")
    parser.add_argument("--font", default="Georgia", help="Nome da fonte principal (padrão: Georgia)")
    parser.add_argument("--size", type=float, default=10.5, help="Tamanho da fonte principal (padrão: 10.5)")
    parser.add_argument("--margin-inner", type=float, default=2.2, help="Margem interna em cm (padrão: 2.2)")
    parser.add_argument("--margin-outer", type=float, default=1.8, help="Margem externa em cm (padrão: 1.8)")
    parser.add_argument("--margin-top", type=float, default=1.8, help="Margem superior em cm (padrão: 1.8)")
    parser.add_argument("--margin-bottom", type=float, default=1.8, help="Margem inferior em cm (padrão: 1.8)")
    parser.add_argument("--title", help="Título customizado para a página de rosto")
    parser.add_argument("--author", help="Autor customizado para a página de rosto")
    parser.add_argument("--no-title-page", action="store_true", help="Ignora a geração da página de rosto")
    parser.add_argument("--no-toc", action="store_true", help="Ignora a geração do sumário")
    parser.add_argument("--keep-docx", action="store_true", help="Mantém o arquivo DOCX temporário gerado")
    
    args = parser.parse_args()
    
    if not os.path.exists(args.epub_path):
        print(f"Erro: Arquivo EPUB não localizado em {args.epub_path}")
        return
        
    print(f"Carregando o arquivo EPUB: {args.epub_path}")
    
    # Gerar caminho temporário para o DOCX
    temp_docx_path = args.pdf_path.rsplit('.', 1)[0] + ".temp.docx"
    
    # 1. Abrir EPUB
    with zipfile.ZipFile(args.epub_path, 'r') as zip_ref:
        title, author, spine_files, opf_dir = get_epub_info(zip_ref)
        
        # Sobrescrever metadados se fornecido
        title = args.title if args.title else title
        author = args.author if args.author else author
        
        print(f"Livro identificado: {title}")
        print(f"Autor identificado: {author}")
        print(f"Total de arquivos de capítulo a processar: {len(spine_files)}")
        
        # 2. Criar documento Word
        document = Document()
        
        # Configurar estilos de fonte padrões
        style_normal = document.styles['Normal']
        style_normal.font.name = args.font
        style_normal.font.size = Pt(args.size)
        style_normal.font.color.rgb = RGBColor(0, 0, 0)
        
        # Configurar estilos de cabeçalho
        h1_style = document.styles['Heading 1']
        h1_style.font.name = args.font
        h1_style.font.size = Pt(16)
        h1_style.font.bold = True
        h1_style.font.color.rgb = RGBColor(120, 0, 0) # Vermelho escuro elegante
        
        h2_style = document.styles['Heading 2']
        h2_style.font.name = args.font
        h2_style.font.size = Pt(13)
        h2_style.font.bold = True
        h2_style.font.color.rgb = RGBColor(0, 0, 0)
        
        h3_style = document.styles['Heading 3']
        h3_style.font.name = args.font
        h3_style.font.size = Pt(11)
        h3_style.font.bold = True
        h3_style.font.italic = True
        h3_style.font.color.rgb = RGBColor(0, 0, 0)
        
        # Margem de impressão para imagens baseada em A5
        max_width_cm = 14.8 - args.margin_inner - args.margin_outer
        
        # 3. Construir Seção 1 (Página de rosto e sumário)
        section_intro = document.sections[0]
        setup_section_layout(section_intro, args.margin_inner, args.margin_outer, args.margin_top, args.margin_bottom)
        
        # Adicionar Página de Rosto
        if not args.no_title_page:
            add_title_page(document, title, author)
            
        # Adicionar Sumário
        if not args.no_toc:
            add_table_of_contents(document)
            
        # 4. Iniciar Seção 2 para o Corpo do Livro
        body_section = document.add_section(WD_SECTION_START.NEW_PAGE)
        setup_section_layout(body_section, args.margin_inner, args.margin_outer, args.margin_top, args.margin_bottom)
        
        # Desvincular cabeçalhos/rodapés da seção anterior
        body_section.header.is_linked_to_previous = False
        body_section.footer.is_linked_to_previous = False
        
        # Adicionar running header com o título do livro
        p_header = body_section.header.paragraphs[0]
        p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_header.paragraph_format.space_after = Pt(12)
        run_header = p_header.add_run(title)
        run_header.font.name = args.font
        run_header.font.size = Pt(8.5)
        run_header.italic = True
        
        # Adicionar rodapé com paginação automática
        p_footer = body_section.footer.paragraphs[0]
        p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_footer = p_footer.add_run()
        run_footer.font.name = args.font
        run_footer.font.size = Pt(9)
        add_page_number(run_footer)
        
        # 5. Processar Capítulos
        just_after_heading = [True]
        is_first_chapter = True
        
        for idx, (file_path, media_type) in enumerate(spine_files):
            print(f"Processando arquivo [{idx + 1}/{len(spine_files)}]: {os.path.basename(file_path)}")
            
            try:
                html_content = zip_ref.read(file_path)
                soup = BeautifulSoup(html_content, 'html.parser')
                body = soup.find('body')
                if not body:
                    continue
                
                # Preprocessar nós de texto órfãos direto no body
                for child in list(body.children):
                    if isinstance(child, str) and child.strip():
                        new_p = soup.new_tag('p')
                        child.replace_with(new_p)
                        new_p.append(child)
                        
                # Adicionar quebra de página entre capítulos
                if not is_first_chapter:
                    document.add_page_break()
                    just_after_heading[0] = True
                else:
                    is_first_chapter = False
                    
                html_dir = os.path.dirname(file_path)
                
                # Processar blocos do body
                for child in body.children:
                    if child.name in ['p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'ul', 'ol', 'table', 'blockquote', 'hr', 'img', 'div', 'section']:
                        process_block(child, document, zip_ref, html_dir, max_width_cm, just_after_heading)
                        
            except Exception as e:
                print(f"Erro ao processar o arquivo {file_path}: {e}")
                
        # 6. Salvar documento DOCX temporário
        print(f"Gravando arquivo intermediário DOCX: {temp_docx_path}")
        document.save(temp_docx_path)
        print("Arquivo intermediário DOCX gravado.")
        
        # 7. Chamar conversão via COM do MS Word
        success = convert_docx_to_pdf(temp_docx_path, args.pdf_path)
        
        # 8. Limpeza do DOCX intermediário
        if not args.keep_docx:
            if os.path.exists(temp_docx_path):
                try:
                    os.remove(temp_docx_path)
                    print("Arquivo DOCX temporário removido.")
                except Exception as e:
                    print(f"Aviso: Não foi possível remover o arquivo DOCX temporário: {e}")
        else:
            print(f"Mantendo o arquivo DOCX temporário em: {temp_docx_path}")
            
        if success:
            print(f"Processo de conversão concluído! PDF gerado com sucesso em: {args.pdf_path}")
        else:
            print("Processo concluído com erros na geração do PDF.")

if __name__ == "__main__":
    main()
